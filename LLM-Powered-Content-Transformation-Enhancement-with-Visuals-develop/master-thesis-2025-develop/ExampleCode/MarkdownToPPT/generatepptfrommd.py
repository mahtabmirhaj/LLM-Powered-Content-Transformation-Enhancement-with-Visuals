import os
import base64
from openai import AzureOpenAI
from dotenv import load_dotenv
from pptx import Presentation

# Environment variable setup
subscription_key = os.getenv('AZURE_OPENAI_API_KEY')
if subscription_key is None:
    # Load the .env file
    print("# Loading from .env file")
    load_dotenv()
    subscription_key = os.getenv('AZURE_OPENAI_API_KEY')
    if subscription_key is None:
        raise ValueError(f"Environment variable AZURE_OPENAI_API_KEY is not set and not found in .env file")
else:
    print("# Env variables already exist")

endpoint = os.getenv('ENDPOINT_URL')
deployment = os.getenv('DEPLOYMENT_NAME')

# Initialize Azure OpenAI Service client with key-based authentication    
client = AzureOpenAI(  
    azure_endpoint=endpoint,  
    api_key=subscription_key,  
    api_version="2024-12-01-preview",
)

def ask_llm(prompt, system_message=None, enforce_json=False):
    """
    Generic function to ask LLM with different prompts
    
    Args:
        prompt: The prompt to send to the LLM
        system_message: Optional system message
        enforce_json: Whether to request JSON response format
    
    Returns:
        The LLM's response as a string
    """
    # Prepare messages
    messages = []
    if system_message:
        messages.append({"role": "system", "content": system_message})
    messages.append({"role": "user", "content": prompt})
    
    # Set params
    params = {
        "model": deployment,
        "messages": messages,
        "reasoning_effort": "high",
        "max_completion_tokens": 4000,
        "stream": False
    }
    
    # Add JSON formatting if requested
    if enforce_json:
        params["response_format"] = {"type": "json_object"}
    
    # Make the API call
    try:
        response = client.chat.completions.create(**params)
        return response.choices[0].message.content
    except Exception as e:
        print(f"Error calling LLM: {e}")
        return None

def generate_ppt_from_markdown(md_path, output_path, slides_config=None):
    """
    Generate a PowerPoint from markdown using LLM prompting at every step
    
    Args:
        md_path: Path to markdown file
        output_path: Path for output PowerPoint
        slides_config: Optional configuration for slides (number of slides, points per slide)
    """
    try:
        # Read the markdown file
        with open(md_path, 'r', encoding='utf-8') as file:
            markdown_content = file.read()
        
        # Step 1: Ask LLM to analyze the markdown and determine optimal slide structure
        if not slides_config:
            analysis_prompt = f"""
            Here's markdown content from a document. Please analyze it and suggest the optimal presentation structure.
            How many slides would best present this information? How many bullet points per slide would be appropriate?
            
            Please format your response as a JSON with 'slide_count' and 'points_per_slide' fields.
            
            The markdown content is:
            {markdown_content[:2000]}... (content truncated for brevity)
            """
            
            print("Analyzing markdown content to determine optimal presentation structure...")
            structure_response = ask_llm(analysis_prompt, 
                                        "You are an expert presentation designer who analyzes content and recommends optimal slide structures.", 
                                        enforce_json=True)
            
            # Parse the response
            import json
            try:
                structure = json.loads(structure_response)
                slide_count = structure.get('slide_count', 10)
                points_per_slide = structure.get('points_per_slide', 7)
                print(f"LLM suggested {slide_count} slides with {points_per_slide} points per slide")
            except:
                # Default if parsing fails
                slide_count = 10
                points_per_slide = 7
                print("Using default structure: 10 slides with 7 points per slide")
        else:
            # Use provided configuration
            slide_count = slides_config.get('slide_count', 10)
            points_per_slide = slides_config.get('points_per_slide', 7)
            
        # Step 2: Ask LLM to generate the presentation content
        content_prompt = f"""
        Based on this markdown content, create a PowerPoint presentation with {slide_count} slides.
        Each slide should have a clear title and {points_per_slide} bullet points.
        
        Focus on extracting the most important information from the markdown.
        Make each bullet point concise, informative, and impactful.
        
        Format your response as a valid JSON object with a 'slides' array.
        Each slide in the array should have a 'title' and 'points' array with {points_per_slide} items.
        
        Here's the markdown content:
        {markdown_content}
        """
        
        print(f"Generating presentation content with {slide_count} slides...")
        content_response = ask_llm(content_prompt, 
                                   "You are an expert at creating effective presentations from documents.", 
                                   enforce_json=True)
        
        # Step 3: Create presentation file using python-pptx
        print("Creating PowerPoint file...")
        import json
        
        # Try to parse the JSON response
        try:
            slides_data = json.loads(content_response)
            
            # Get slides array - handle different possible structures
            if isinstance(slides_data, list):
                slides = slides_data
            else:
                slides = slides_data.get('slides', [])
                
                # If still no slides, check for any array in the response
                if not slides and isinstance(slides_data, dict):
                    for key, value in slides_data.items():
                        if isinstance(value, list) and len(value) > 0:
                            slides = value
                            break
            
            if not slides:
                raise ValueError("No slides found in LLM response")
            
            # Step 4: Ask LLM for design advice
            design_prompt = f"""
            I'm creating a PowerPoint presentation with {slide_count} slides based on the following content.
            Please suggest a cohesive design approach including:
            1. Color scheme (2-3 colors)
            2. Font style suggestions
            3. Any layout recommendations
            
            Return as JSON with 'colors', 'fonts', and 'layout_tips' fields.
            
            Here's a sample of the content:
            {json.dumps(slides[:2], indent=2)}
            """
            
            design_response = ask_llm(design_prompt, 
                                     "You are an expert presentation designer with a strong sense of aesthetics.", 
                                     enforce_json=True)
            
            # Parse design suggestions
            try:
                design = json.loads(design_response)
                print("Design suggestions received from LLM")
            except:
                design = {}
                print("Using default presentation design")
            
            # Step 5: Create and format the presentation
            prs = Presentation()
            
            # Add slides
            for slide_data in slides:
                # Handle different possible structures
                if isinstance(slide_data, str):
                    title = slide_data
                    points = [""] * points_per_slide
                else:
                    # Extract title and points
                    title = slide_data.get('title', 'Slide')
                    points = slide_data.get('points', [])
                    
                    # Alternative key names that might be used
                    if not points and 'content' in slide_data:
                        points = slide_data['content']
                    elif not points and 'bullets' in slide_data:
                        points = slide_data['bullets']
                    
                    # If points is a string, convert to list
                    if isinstance(points, str):
                        points = points.split('\n')
                
                # Ensure we have exactly the right number of points
                while len(points) < points_per_slide:
                    points.append("")
                
                # Use title and content slide layout
                slide_layout = prs.slide_layouts[1]
                slide = prs.slides.add_slide(slide_layout)
                
                # Set title
                title_shape = slide.shapes.title
                title_shape.text = title
                
                # Set content
                content_shape = slide.placeholders[1]
                text_frame = content_shape.text_frame
                
                # Add bullet points
                p = text_frame.paragraphs[0]
                p.text = points[0]
                
                for point in points[1:points_per_slide]:  # Limit to exactly points_per_slide
                    p = text_frame.add_paragraph()
                    p.text = point
                    p.level = 0
            
            # Save the presentation
            prs.save(output_path)
            print(f"PowerPoint presentation created successfully at {output_path}")
            return True
            
        except Exception as e:
            print(f"Error creating PowerPoint: {e}")
            
            # Ask LLM for help with the error
            error_prompt = f"""
            I encountered this error while trying to create a PowerPoint presentation:
            {str(e)}
            
            The LLM response was:
            {content_response[:500]}... (truncated)
            
            What might be causing this issue and how can I fix it?
            """
            
            help_response = ask_llm(error_prompt, "You are an expert at debugging presentation generation issues.")
            print("\nLLM debug suggestion:")
            print(help_response)
            
            return False
    
    except Exception as e:
        print(f"Error generating PowerPoint: {e}")
        return False

# Example usage
if __name__ == "__main__":
    # Input parameters
    md_path = input("Enter the path to the markdown file: ")
    
    # Verify file exists
    if not os.path.exists(md_path):
        print(f"Error: File '{md_path}' not found.")
        exit(1)
    
    default_output = os.path.splitext(os.path.basename(md_path))[0] + ".pptx"
    output_path = input(f"Enter the path for the output PowerPoint file (default: {default_output}): ") or default_output
    
    # Ask if user wants custom slide configuration
    custom_config_option = input("Do you want to specify custom slide configuration? (y/n): ").lower()
    
    slides_config = None
    if custom_config_option == 'y':
        try:
            slide_count = int(input("How many slides would you like? (default: 10): ") or "10")
            points_per_slide = int(input("How many bullet points per slide? (default: 7): ") or "7")
            slides_config = {
                'slide_count': slide_count,
                'points_per_slide': points_per_slide
            }
        except ValueError:
            print("Invalid input. Using default configuration.")
    
    print("\nGenerating PowerPoint presentation from markdown...")
    success = generate_ppt_from_markdown(md_path, output_path, slides_config)
    
    if success:
        print(f"\nSuccess! PowerPoint created at {output_path}")
    else:
        print("\nFailed to create PowerPoint. Check the errors above for details.")