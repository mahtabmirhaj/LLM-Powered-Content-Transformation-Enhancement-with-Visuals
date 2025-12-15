from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from typing import List, Dict
import json

class PresentationCreator:
    def __init__(self, text_analyzer):
        self.text_analyzer = text_analyzer

    def create_presentation(self, title: str, summary: str, key_points: List[str], output_path: str):
        """Create an enhanced PowerPoint presentation using LLM-generated content."""
        try:
            # Get presentation content from text analyzer
            presentation_content = self.text_analyzer.generate_presentation_content(
                title, summary, key_points
            )
            
            prs = Presentation()
            
            # Create title slide
            self._create_title_slide(prs, presentation_content["title_slide"])
            
            # Create summary slide
            self._create_summary_slide(prs, presentation_content["summary_slide"])
            
            # Create content slides
            for slide_content in presentation_content["content_slides"]:
                self._create_content_slide(prs, slide_content)
            
            try:
                prs.save(output_path)
            except Exception as e:
                raise Exception(f"Error saving presentation: {str(e)}")

        except Exception as e:
            raise Exception(f"Error creating presentation: {str(e)}")

    def _create_title_slide(self, prs, content):
        """Create an engaging title slide."""
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        title = slide.shapes.title
        title.text = content["title"]
        
        # Enhance formatting
        title_tf = title.text_frame
        title_tf.paragraphs[0].alignment = PP_ALIGN.CENTER
        title_tf.paragraphs[0].font.size = Pt(44)

    def _create_summary_slide(self, prs, content):
        """Create a well-formatted summary slide."""
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        
        title = slide.shapes.title
        title.text = content["title"]
        
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.text = content["content"]
        
        # Enhance formatting
        for paragraph in tf.paragraphs:
            paragraph.font.size = Pt(18)

    def _create_content_slide(self, prs, content):
        """Create a content slide with enhanced formatting."""
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        
        title = slide.shapes.title
        title.text = content["title"]
        
        body = slide.placeholders[1]
        tf = body.text_frame
        tf.clear()  # Clear existing text
        
        for point in content["points"]:
            p = tf.add_paragraph()
            p.text = f"• {point}"
            p.level = 0
            p.font.size = Pt(18)
            
        # Enhance formatting
        title_tf = title.text_frame
        title_tf.paragraphs[0].font.size = Pt(32)